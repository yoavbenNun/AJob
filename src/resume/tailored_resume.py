"""Human-approved resume proposals and professional document generation."""

from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path
import os
import re
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.matching.resume_matcher import MatchResult, match_resume_to_job
from src.models import JobPosting, ResumeProfile


@dataclass(frozen=True, slots=True)
class TailoringChange:
    """One factual, independently approvable resume change."""

    id: str
    section: str
    description: str
    original: str
    proposed: str
    evidence: str
    score_delta: int = 0


@dataclass(frozen=True, slots=True)
class TailoringProposal:
    """A reviewable proposal before any separate resume file is created."""

    before_score: int
    after_score: int
    changes: tuple[TailoringChange, ...]
    safeguards: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class ResumeArtifacts:
    """The two application-ready files generated from approved content."""

    docx_path: Path
    pdf_path: Path


def build_tailoring_proposal(
    profile: ResumeProfile, job: JobPosting, match: MatchResult
) -> TailoringProposal:
    """Build deterministic changes whose claims are supported by the resume."""
    changes: list[TailoringChange] = [
        TailoringChange(
            id="professional_format",
            section="Layout",
            description="Use a clean, single-column resume layout for DOCX and PDF.",
            original="Formatting from the source resume",
            proposed="Consistent typography, spacing, headings, and application-ready file names",
            evidence="This changes presentation only and does not alter resume claims.",
        )
    ]
    ordered_skills = _ordered_skills(profile, match)
    if ordered_skills and ordered_skills != profile.skills:
        changes.append(
            TailoringChange(
                id="prioritize_matched_skills",
                section="Technical Skills",
                description="Move job-matched skills to the beginning of the skills list.",
                original=", ".join(profile.skills),
                proposed=", ".join(ordered_skills),
                evidence=f"Matched by the deterministic scorer: {', '.join(match.matched_skills)}.",
            )
        )
    summary = _targeted_summary(profile, job, match)
    if summary and summary != profile.summary:
        changes.append(
            TailoringChange(
                id="targeted_summary",
                section="Summary",
                description="Add a concise role-specific summary using documented skills only.",
                original=profile.summary or "No summary in the source resume",
                proposed=summary,
                evidence=(
                    "Every named skill is present in the parsed source resume: "
                    f"{', '.join(match.matched_skills)}."
                ),
            )
        )
    after_profile = apply_tailoring_changes(
        profile, match, tuple(change.id for change in changes), job=job
    )
    after_score = match_resume_to_job(after_profile, job).score
    return TailoringProposal(
        before_score=match.score,
        after_score=after_score,
        changes=tuple(changes),
        safeguards=(
            "Only selected changes are applied.",
            "No new skills, experience, projects, metrics, or credentials are invented.",
            "The master resume remains unchanged; separate DOCX and PDF files are created.",
        ),
    )


def apply_tailoring_changes(
    profile: ResumeProfile,
    match: MatchResult,
    approved_change_ids: tuple[str, ...],
    *,
    job: JobPosting | None = None,
) -> ResumeProfile:
    """Apply only recognized change IDs to a copy of the parsed profile."""
    selected = set(approved_change_ids)
    known = {"professional_format", "prioritize_matched_skills", "targeted_summary"}
    unknown = selected - known
    if unknown:
        raise ValueError(f"Unknown tailoring change: {', '.join(sorted(unknown))}")
    if not selected:
        raise PermissionError("Select at least one tailoring change before generation.")
    skills = (
        _ordered_skills(profile, match)
        if "prioritize_matched_skills" in selected
        else profile.skills
    )
    summary = profile.summary
    if "targeted_summary" in selected:
        summary = _targeted_summary_from_match(profile, match, job)
    return replace(profile, skills=skills, summary=summary)


def tailored_resume_filename(job: JobPosting, suffix: str = ".docx") -> str:
    """Return a safe, readable filename for the selected company and role."""
    if suffix.casefold() not in {".docx", ".pdf"}:
        raise ValueError("Tailored resume suffix must be .docx or .pdf.")
    stem = f"{_filename_component(job.company, 'company')}-{_filename_component(job.title, 'role')}-tailored"
    return f"{stem[:180]}{suffix.casefold()}"


def tailored_resume_output_path(output_directory: str | Path, job: JobPosting) -> Path:
    return Path(output_directory) / tailored_resume_filename(job)


def create_tailored_resume(
    profile: ResumeProfile,
    job: JobPosting,
    match: MatchResult,
    output_path: str | Path,
    *,
    approved: bool,
    approved_change_ids: tuple[str, ...] = ("professional_format", "prioritize_matched_skills"),
) -> Path:
    """Compatibility wrapper that creates one approved DOCX."""
    if not approved:
        raise PermissionError("Tailored resume generation requires explicit approval.")
    target = Path(output_path)
    _validate_target(profile, target, ".docx")
    tailored = apply_tailoring_changes(profile, match, approved_change_ids, job=job)
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=target.parent, prefix=".resume-") as temp_dir:
        temporary = Path(temp_dir) / target.name
        _write_docx(tailored, temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise ValueError("Generated DOCX is empty.")
        Document(temporary)
        os.link(temporary, target)
    return target


def create_resume_artifacts(
    profile: ResumeProfile,
    job: JobPosting,
    match: MatchResult,
    output_directory: str | Path,
    *,
    approved: bool,
    approved_change_ids: tuple[str, ...],
) -> ResumeArtifacts:
    """Atomically publish matching DOCX and PDF files after explicit approval."""
    if not approved:
        raise PermissionError("Resume artifact generation requires explicit approval.")
    tailored = apply_tailoring_changes(profile, match, approved_change_ids, job=job)
    output_root = Path(output_directory)
    docx_target = output_root / tailored_resume_filename(job, ".docx")
    pdf_target = output_root / tailored_resume_filename(job, ".pdf")
    _validate_target(profile, docx_target, ".docx")
    _validate_target(profile, pdf_target, ".pdf")
    output_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output_root, prefix=".resume-") as temp_dir:
        temp_root = Path(temp_dir)
        temp_docx = temp_root / docx_target.name
        temp_pdf = temp_root / pdf_target.name
        _write_docx(tailored, temp_docx)
        _write_pdf(tailored, temp_pdf)
        _verify_artifacts(tailored, temp_docx, temp_pdf)
        os.link(temp_docx, docx_target)
        try:
            os.link(temp_pdf, pdf_target)
        except Exception:
            if docx_target.exists() and os.path.samefile(temp_docx, docx_target):
                docx_target.unlink()
            raise
    return ResumeArtifacts(docx_path=docx_target, pdf_path=pdf_target)


def _write_docx(profile: ResumeProfile, target: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.46)
    section.bottom_margin = Inches(0.46)
    section.left_margin = Inches(0.64)
    section.right_margin = Inches(0.64)
    _configure_docx_styles(document)

    name = document.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.add_run(profile.name or "Resume")
    contact_values = [value for value in (profile.location, profile.email, profile.phone) if value]
    if contact_values:
        contact = document.add_paragraph(" | ".join(contact_values))
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.style = document.styles["Contact"]
    profile_links, project_links = _split_profile_links(profile.links)
    if profile_links:
        links_paragraph = document.add_paragraph()
        links_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        links_paragraph.style = document.styles["Contact"]
        for index, (label, url) in enumerate(profile_links):
            if index:
                links_paragraph.add_run(" | ")
            _add_hyperlink(links_paragraph, label, url)
    _add_docx_section(document, "Summary", profile.summary)
    _add_docx_section(document, "Technical Skills", ", ".join(profile.skills))
    _add_docx_section(document, "Experience", profile.experience)
    _add_docx_section(document, "Projects", profile.projects, links=project_links)
    _add_docx_section(document, "Education", profile.education)
    document.save(target)


def _configure_docx_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(1)
    normal.paragraph_format.line_spacing = 1.0
    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(2)
    heading = styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(10.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(1)
    contact = styles.add_style("Contact", 1)
    contact.font.name = "Arial"
    contact.font.size = Pt(8.5)
    contact.font.color.rgb = RGBColor(65, 65, 65)
    contact.paragraph_format.space_after = Pt(4)


def _add_docx_section(
    document: Document, heading: str, content: str, links: tuple[str, ...] = ()
) -> None:
    if not content.strip():
        return
    document.add_heading(heading, level=1)
    link_index = 0
    for line in _content_lines(content):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        if _looks_like_bullet(line):
            paragraph.style = document.styles["List Bullet"]
            line = line.lstrip("-*• ")
        if line.casefold() == "git link" and link_index < len(links):
            _add_hyperlink(paragraph, "GitHub repository", links[link_index])
            link_index += 1
        else:
            paragraph.add_run(line.lstrip("| "))
    for url in links[link_index:]:
        paragraph = document.add_paragraph()
        _add_hyperlink(paragraph, "Project link", url)


def _write_pdf(profile: ResumeProfile, target: Path) -> None:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    resume_text = "\n".join(
        (
            profile.name, profile.location, profile.summary, profile.experience,
            profile.projects, profile.education, *profile.skills,
        )
    )
    regular_font, bold_font = _pdf_font_paths(
        require_hebrew=bool(re.search(r"[\u0590-\u05ff]", resume_text))
    )
    pdfmetrics.registerFont(TTFont("ResumeSans", regular_font))
    pdfmetrics.registerFont(TTFont("ResumeSansBold", bold_font))
    name_style = ParagraphStyle(
        "ResumeName", parent=styles["Title"], fontName="ResumeSansBold",
        fontSize=18, leading=19, alignment=TA_CENTER, textColor="#000000",
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "Contact", parent=styles["Normal"], fontName="ResumeSans", fontSize=8.8,
        leading=10, alignment=TA_CENTER, textColor="#414141", spaceAfter=4,
    )
    heading_style = ParagraphStyle(
        "Section", parent=styles["Heading1"], fontName="ResumeSansBold",
        fontSize=10.3, leading=11.5, textColor="#000000", spaceBefore=4, spaceAfter=1,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["BodyText"], fontName="ResumeSans", fontSize=9.2,
        leading=10.3, textColor="#000000", spaceAfter=0.6,
    )
    bullet_style = ParagraphStyle(
        "Bullet", parent=body_style, leftIndent=10, firstLineIndent=-7,
    )
    story = [Paragraph(escape(_pdf_display(profile.name or "Resume")), name_style)]
    contact = " | ".join(value for value in (profile.location, profile.email, profile.phone) if value)
    if contact:
        story.append(Paragraph(escape(_pdf_display(contact)), contact_style))
    profile_links, project_links = _split_profile_links(profile.links)
    if profile_links:
        link_markup = " | ".join(
            f'<link href="{escape(url)}" color="#414141">{escape(label)}</link>'
            for label, url in profile_links
        )
        story.append(Paragraph(link_markup, contact_style))
    for heading, content in (
        ("Summary", profile.summary),
        ("Technical Skills", ", ".join(profile.skills)),
        ("Experience", profile.experience),
        ("Projects", profile.projects),
        ("Education", profile.education),
    ):
        if not content.strip():
            continue
        story.append(Paragraph(escape(_pdf_display(heading)), heading_style))
        project_link_index = 0
        for line in _content_lines(content):
            is_bullet = _looks_like_bullet(line)
            clean = line.lstrip("-*• ") if is_bullet else line
            prefix = "• " if is_bullet else ""
            if heading == "Projects" and clean.casefold() == "git link" and project_link_index < len(project_links):
                url = project_links[project_link_index]
                project_link_index += 1
                markup = f'<link href="{escape(url)}" color="#1f4e79">GitHub repository</link>'
            else:
                markup = prefix + escape(_pdf_display(clean.lstrip("| ")))
            story.append(Paragraph(markup, bullet_style if is_bullet else body_style))
        if heading == "Projects":
            for url in project_links[project_link_index:]:
                story.append(Paragraph(
                    f'<link href="{escape(url)}" color="#1f4e79">Project link</link>',
                    body_style,
                ))
    document = SimpleDocTemplate(
        str(target), pagesize=letter, rightMargin=0.64 * inch, leftMargin=0.64 * inch,
        topMargin=0.46 * inch, bottomMargin=0.46 * inch,
        title=f"{profile.name or 'Candidate'} Resume", author=profile.name or "Candidate",
    )
    document.build(story)


def _verify_artifacts(profile: ResumeProfile, docx_path: Path, pdf_path: Path) -> None:
    if not docx_path.is_file() or docx_path.stat().st_size == 0:
        raise ValueError("Generated DOCX is empty.")
    if not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise ValueError("Generated PDF is empty.")
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    from pypdf import PdfReader
    reader = PdfReader(pdf_path)
    if not reader.pages:
        raise ValueError("Generated PDF has no pages.")
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for expected in (profile.name, *profile.skills):
        pdf_expected = _pdf_display(expected)
        if expected and (
            expected not in docx_text
            or (expected not in pdf_text and pdf_expected not in pdf_text)
        ):
            raise ValueError(f"Generated artifacts are missing approved content: {expected}")


def _validate_target(profile: ResumeProfile, target: Path, suffix: str) -> None:
    if target.suffix.casefold() != suffix:
        raise ValueError(f"Tailored resume output must use the {suffix} extension.")
    source = Path(profile.source_path)
    if profile.source_path != "<memory>" and source.resolve() == target.resolve():
        raise ValueError("Tailored resume output must be separate from the master resume.")
    if target.exists():
        raise FileExistsError(f"Refusing to overwrite existing file: {target}")


def _ordered_skills(profile: ResumeProfile, match: MatchResult) -> tuple[str, ...]:
    return tuple(match.matched_skills) + tuple(
        skill for skill in profile.skills if skill not in match.matched_skills
    )


def _targeted_summary(profile: ResumeProfile, job: JobPosting, match: MatchResult) -> str:
    return _targeted_summary_from_match(profile, match, job)


def _targeted_summary_from_match(
    profile: ResumeProfile, match: MatchResult, job: JobPosting | None = None
) -> str:
    if not match.matched_skills:
        return profile.summary
    skills = ", ".join(match.matched_skills[:5])
    base = f"Software engineering candidate with documented skills in {skills}."
    if job is not None:
        base += f" Interested in the {job.title} role at {job.company}."
    return base


def _looks_like_bullet(value: str) -> bool:
    return value.startswith(("-", "*", "•"))


def _content_lines(content: str) -> tuple[str, ...]:
    """Remove extraction-only separator lines without rewriting resume text."""
    return tuple(
        line for line in (part.strip() for part in content.splitlines())
        if line and line not in {"|", "•"}
    )


def _split_profile_links(
    links: tuple[str, ...]
) -> tuple[tuple[tuple[str, str], ...], tuple[str, ...]]:
    profile_links: list[tuple[str, str]] = []
    project_links: list[str] = []
    for url in links:
        lowered = url.casefold()
        if "linkedin.com/" in lowered and not any(label == "LinkedIn" for label, _ in profile_links):
            profile_links.append(("LinkedIn", url))
        elif "github.com/" in lowered and url.rstrip("/").count("/") <= 3 and not any(label == "GitHub" for label, _ in profile_links):
            profile_links.append(("GitHub", url))
        else:
            project_links.append(url)
    return tuple(profile_links), tuple(project_links)


def _add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _pdf_font_paths(*, require_hebrew: bool) -> tuple[Path, Path]:
    reportlab_fonts = Path(__import__("reportlab").__file__).parent / "fonts"
    unicode_candidates = (
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    )
    fallback = ((reportlab_fonts / "Vera.ttf", reportlab_fonts / "VeraBd.ttf"),)
    candidates = unicode_candidates if require_hebrew else (*unicode_candidates, *fallback)
    try:
        return next(
            (regular, bold) for regular, bold in candidates
            if regular.is_file() and bold.is_file()
        )
    except StopIteration as error:
        language = "Hebrew-capable " if require_hebrew else ""
        raise RuntimeError(f"No {language}PDF font is available.") from error


def _pdf_display(value: str) -> str:
    if not re.search(r"[\u0590-\u05ff]", value):
        return value
    from bidi.algorithm import get_display

    return get_display(value)


def _filename_component(value: str, fallback: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "-", value)
    cleaned = re.sub(r"\s+", "-", cleaned.strip()).strip(".-")
    return cleaned or fallback


__all__ = [
    "ResumeArtifacts", "TailoringChange", "TailoringProposal",
    "apply_tailoring_changes", "build_tailoring_proposal",
    "create_resume_artifacts", "create_tailored_resume",
    "tailored_resume_filename", "tailored_resume_output_path",
]
