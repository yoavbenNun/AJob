"""Local resume text extraction and deterministic profile parsing."""

import re
from pathlib import Path

from src.models import ResumeProfile


SECTION_ALIASES = {
    "summary": {"summary", "profile", "professional summary", "objective", "about"},
    "skills": {"skills", "technical skills", "technologies", "tech stack"},
    "experience": {
        "experience", "work experience", "professional experience",
        "employment", "employment history",
    },
    "education": {"education", "academic background", "academic experience"},
    "projects": {"projects", "personal projects", "academic projects", "selected projects"},
}

SKILL_ALIASES = {
    "Python": ("python",), "Java": ("java",), "C": ("c",),
    "C++": ("c++", "cpp"), "C#": ("c#", "c sharp"),
    "JavaScript": ("javascript", "js"), "TypeScript": ("typescript", "ts"),
    "SQL": ("sql",), "HTML": ("html", "html5"), "CSS": ("css", "css3"),
    "React": ("react", "react.js", "reactjs"), "Node.js": ("node.js", "nodejs"),
    "FastAPI": ("fastapi",), "Flask": ("flask",), "Django": ("django",),
    "Spring Boot": ("spring boot",), "PostgreSQL": ("postgresql", "postgres"),
    "MySQL": ("mysql",), "MongoDB": ("mongodb",), "SQLite": ("sqlite",),
    "Redis": ("redis",), "Docker": ("docker",),
    "Kubernetes": ("kubernetes", "k8s"),
    "AWS": ("aws", "amazon web services"), "Azure": ("azure",),
    "Google Cloud": ("google cloud", "gcp"), "Git": ("git",),
    "GitHub Actions": ("github actions",), "Linux": ("linux",),
    "REST APIs": ("rest api", "restful api", "rest apis"),
    "GraphQL": ("graphql",), "Playwright": ("playwright",),
    "Pytest": ("pytest",), "Machine Learning": ("machine learning",),
    "Pandas": ("pandas",), "NumPy": ("numpy",),
    "TensorFlow": ("tensorflow",), "PyTorch": ("pytorch",),
    "Qt/QML": ("qt", "qt6", "qml"), "CMake": ("cmake",),
    "WebSockets": ("websocket", "websockets"),
    "Express.js": ("express.js", "expressjs"),
    "Algorithms": ("algorithm", "algorithms"),
    "Data Structures": ("data structure", "data structures"),
    "Object-Oriented Programming": ("object-oriented programming", "oop"),
    "Multithreading": ("multithreading", "multi-threading"),
    "Microservices": ("microservice", "microservices"),
    "Unit Testing": ("unit test", "unit tests", "unit testing"),
    "Integration Testing": ("integration test", "integration tests", "integration testing"),
    "Regression Testing": ("regression test", "regression tests", "regression testing"),
    "CTest": ("ctest",),
}


def parse_resume(path: str | Path) -> ResumeProfile:
    """Read a supported resume file and return a normalized profile."""
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"Resume file not found: {source}")
    raw_text = extract_resume_text(source)
    if not raw_text.strip():
        raise ValueError(f"No readable text found in resume: {source}")
    return parse_resume_text(raw_text, source_path=str(source))


def extract_resume_text(path: str | Path) -> str:
    """Extract text from TXT, Markdown, PDF, or DOCX without external services."""
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix in {".txt", ".md"}:
        return source.read_text(encoding="utf-8-sig")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(source)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(source)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join((*paragraphs, *cells))
    raise ValueError(
        f"Unsupported resume format '{source.suffix}'. Use .txt, .md, .pdf, or .docx."
    )


def parse_resume_text(text: str, source_path: str = "<memory>") -> ResumeProfile:
    """Parse contact details, sections, and normalized skills from resume text."""
    normalized = _normalize_text(text)
    sections = split_sections(normalized)
    header = sections.get("header", "")
    email_match = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", normalized)
    phone_match = re.search(r"(?<!\w)(?:\+?\d[\d ()-]{7,}\d)(?!\w)", header)
    return ResumeProfile(
        source_path=source_path,
        name=_extract_name(header),
        location=_extract_location(header),
        email=email_match.group(0) if email_match else "",
        phone=" ".join(phone_match.group(0).split()) if phone_match else "",
        skills=extract_skills(normalized),
        education=sections.get("education", ""),
        experience=sections.get("experience", ""),
        projects=sections.get("projects", ""),
        summary=sections.get("summary", ""),
        raw_text=normalized,
    )


def split_sections(text: str) -> dict[str, str]:
    """Split resume text using common standalone section headings."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.splitlines():
        heading = _canonical_heading(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return {
        name: "\n".join(lines).strip()
        for name, lines in sections.items()
        if any(line.strip() for line in lines)
    }


def extract_skills(text: str) -> tuple[str, ...]:
    """Return known skills in stable canonical order without substring matches."""
    lowered = text.casefold()
    return tuple(
        canonical
        for canonical, aliases in SKILL_ALIASES.items()
        if any(_contains_skill(lowered, alias.casefold()) for alias in aliases)
    )


def _contains_skill(text: str, alias: str) -> bool:
    return re.search(rf"(?<![\w+#]){re.escape(alias)}(?![\w+#])", text) is not None


def _canonical_heading(line: str) -> str | None:
    candidate = re.sub(r"[:|]+$", "", line.strip()).casefold()
    for canonical, aliases in SECTION_ALIASES.items():
        if candidate in aliases:
            return canonical
    return None


def _extract_name(header: str) -> str:
    lines = [line.strip() for line in header.splitlines() if line.strip()]
    for index, candidate in enumerate(lines):
        if "@" in candidate or any(char.isdigit() for char in candidate):
            continue
        if (
            index + 1 < len(lines)
            and candidate.isupper()
            and _looks_like_name_fragment(candidate)
            and lines[index + 1].isupper()
            and _looks_like_name_fragment(lines[index + 1])
        ):
            return f"{candidate} {lines[index + 1]}".title()
        return candidate
    return ""


def _looks_like_name_fragment(value: str) -> bool:
    role_words = {"developer", "engineer", "engineering", "student", "manager"}
    words = value.casefold().replace("-", " ").replace("'", " ").split()
    return 1 <= len(words) <= 3 and not role_words.intersection(words) and all(
        word.isalpha() for word in words
    )


def _extract_location(header: str) -> str:
    """Extract an explicit city/country line from the resume header only."""
    for line in header.splitlines():
        candidate = " ".join(line.strip().split())
        if re.fullmatch(r"[A-Za-z][A-Za-z .'-]+,\s*[A-Za-z][A-Za-z .'-]+", candidate):
            if "@" not in candidate and not any(char.isdigit() for char in candidate):
                return candidate
    return ""


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()
