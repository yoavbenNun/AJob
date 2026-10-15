from pathlib import Path
from dataclasses import replace

import pytest
from docx import Document

from src.matching.resume_matcher import match_resume_to_job
from src.models import JobPosting
from src.resume.resume_parser import parse_resume_text
from src.resume.tailored_resume import (
    apply_tailoring_changes,
    build_tailoring_proposal,
    create_resume_artifacts,
    create_tailored_resume,
    tailored_resume_filename,
)


def _inputs():
    profile = parse_resume_text(
        "Dana Developer\ndana@example.com\nSKILLS\nPython, Git, Docker\n"
        "EXPERIENCE\nBuilt a Python service.\nEDUCATION\nB.Sc. Computer Science"
    )
    profile = replace(profile, skills=("Git", "Python", "Docker"))
    job = JobPosting(
        "Python Intern", "Example", "Remote", "https://example.com/jobs/1",
        requirements="Python and Docker",
    )
    return profile, job, match_resume_to_job(profile, job)


def test_tailored_resume_requires_explicit_approval(tmp_path: Path) -> None:
    profile, job, match = _inputs()

    with pytest.raises(PermissionError, match="explicit approval"):
        create_tailored_resume(profile, job, match, tmp_path / "resume.docx", approved=False)


def test_tailored_resume_uses_only_existing_content_and_reorders_skills(tmp_path: Path) -> None:
    profile, job, match = _inputs()
    output = tmp_path / "resume.docx"

    create_tailored_resume(profile, job, match, output, approved=True)

    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "Dana Developer" in text
    assert "Python, Docker, Git" in text
    assert "Built a Python service." in text
    assert "Kubernetes" not in text


def test_tailored_resume_refuses_to_overwrite(tmp_path: Path) -> None:
    profile, job, match = _inputs()
    output = tmp_path / "resume.docx"
    output.write_bytes(b"existing")

    with pytest.raises(FileExistsError):
        create_tailored_resume(profile, job, match, output, approved=True)


def test_tailoring_proposal_keeps_score_truthful_and_names_job_output() -> None:
    profile, job, match = _inputs()

    proposal = build_tailoring_proposal(profile, job, match)

    assert proposal.before_score == proposal.after_score == match.score
    changes = {change.id: change for change in proposal.changes}
    assert changes["prioritize_matched_skills"].proposed == "Python, Docker, Git"
    assert changes["targeted_summary"].evidence.startswith("Every named skill")
    assert any("No new skills" in safeguard for safeguard in proposal.safeguards)
    assert tailored_resume_filename(job) == "Example-Python-Intern-tailored.docx"


def test_tailoring_keeps_score_unchanged_when_it_only_changes_presentation() -> None:
    profile = parse_resume_text(
        "Dana Developer\nSKILLS\nPython\nEXPERIENCE\n"
        "Used Kubernetes to deploy a service.\nEDUCATION\nB.Sc. Computer Science"
    )
    job = JobPosting(
        "Software Intern", "Example", "Remote", "https://example.com/jobs/2",
        requirements="Python and Kubernetes",
    )
    match = match_resume_to_job(profile, job)

    proposal = build_tailoring_proposal(profile, job, match)
    assert "Kubernetes" in profile.skills
    assert proposal.before_score == proposal.after_score == match.score
    assert all(change.score_delta == 0 for change in proposal.changes)


def test_tailored_resume_cannot_replace_master_resume(tmp_path: Path) -> None:
    profile, job, match = _inputs()
    master = tmp_path / "master.docx"
    profile = profile.__class__(
        source_path=str(master), name=profile.name, location=profile.location,
        email=profile.email, phone=profile.phone, skills=profile.skills,
        education=profile.education, experience=profile.experience,
        projects=profile.projects, summary=profile.summary, raw_text=profile.raw_text,
    )

    with pytest.raises(ValueError, match="separate from the master"):
        create_tailored_resume(profile, job, match, master, approved=True)


def test_apply_tailoring_changes_rejects_empty_and_unknown_approval() -> None:
    profile, job, match = _inputs()

    with pytest.raises(PermissionError, match="at least one"):
        apply_tailoring_changes(profile, match, ())
    with pytest.raises(ValueError, match="Unknown"):
        apply_tailoring_changes(profile, match, ("invent_experience",))


def test_create_resume_artifacts_writes_matching_docx_and_pdf(tmp_path: Path) -> None:
    profile, job, match = _inputs()
    profile = replace(
        profile,
        links=(
            "https://www.linkedin.com/in/dana/",
            "https://github.com/dana",
            "https://github.com/dana/python-service",
            "https://portfolio.example/dana-project",
        ),
        projects="Python Service\nGit Link\nBuilt a Python service.",
    )

    artifacts = create_resume_artifacts(
        profile,
        job,
        match,
        tmp_path,
        approved=True,
        approved_change_ids=("professional_format", "prioritize_matched_skills"),
    )

    assert artifacts.docx_path.is_file()
    assert artifacts.pdf_path.is_file()
    docx_text = "\n".join(p.text for p in Document(artifacts.docx_path).paragraphs)
    assert "Python, Docker, Git" in docx_text
    assert "Kubernetes" not in docx_text
    docx_links = {
        relationship.target_ref
        for relationship in Document(artifacts.docx_path).part.rels.values()
        if relationship.reltype.endswith("/hyperlink")
    }
    assert set(profile.links) <= docx_links
    from pypdf import PdfReader
    pdf_links = {
        annotation.get_object().get("/A", {}).get("/URI")
        for page in PdfReader(artifacts.pdf_path).pages
        for annotation in (page.get("/Annots") or ())
    }
    assert set(profile.links) <= pdf_links


def test_create_resume_artifacts_preflights_both_outputs(tmp_path: Path) -> None:
    profile, job, match = _inputs()
    existing_pdf = tmp_path / "Example-Python-Intern-tailored.pdf"
    existing_pdf.write_bytes(b"existing")

    with pytest.raises(FileExistsError):
        create_resume_artifacts(
            profile, job, match, tmp_path, approved=True,
            approved_change_ids=("professional_format",),
        )

    assert not (tmp_path / "Example-Python-Intern-tailored.docx").exists()


def test_targeted_summary_does_not_infer_study_field_from_dates() -> None:
    profile, job, match = _inputs()
    profile = replace(profile, education="History | Example University\n2024-2027")

    proposal = build_tailoring_proposal(profile, job, match)
    summary = next(change for change in proposal.changes if change.id == "targeted_summary")

    assert summary.proposed.startswith("Software engineering candidate")
    assert "Software Engineering student" not in summary.proposed


def test_pdf_generation_supports_hebrew_contact_text(tmp_path: Path) -> None:
    profile, job, match = _inputs()
    profile = replace(profile, name="יואב בן נון", location="ראשון לציון, ישראל")
    assert any("\u0590" <= character <= "\u05ff" for character in profile.name)

    artifacts = create_resume_artifacts(
        profile, job, match, tmp_path, approved=True,
        approved_change_ids=("professional_format",),
    )

    from pypdf import PdfReader
    text = "\n".join(page.extract_text() or "" for page in PdfReader(artifacts.pdf_path).pages)
    assert "ןונ ןב באוי" in text or "יואב בן נון" in text
