from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

from docx import Document
from fastapi.testclient import TestClient
import pytest

from src.models import JobPosting
from src.resume.resume_manager import save_active_resume
from src.storage.job_repository import JobRepository
from src.web import app as web


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.delenv('DATABASE_URL', raising=False)
    repo = JobRepository(tmp_path / 'jobs.db')
    repo.initialize()
    ids = [repo.save(JobPosting(title, 'Demo Systems', 'Remote, Israel',
        f'https://example.test/jobs/{i}', description='Develop Python software services.',
        requirements='Currently pursuing a BSc in Software Engineering. Python and Git.', status=status))
        for i, (title, status) in enumerate([
            ('Software Student', 'new'), ('Software Intern', 'reviewed'), ('Junior Software Engineer', 'applied')])]
    doc = Document()
    for line in ['Demo Candidate', 'SKILLS', 'Python, Git', 'PROJECTS', 'Built Python services with Git.',
                 'EDUCATION', 'B.Sc. Software Engineering (in progress), third-year, expected 2027']:
        doc.add_paragraph(line)
    content = BytesIO()
    doc.save(content)
    save_active_resume(tmp_path, 'demo.docx', content.getvalue())
    app = web.create_app(tmp_path / 'jobs.db', resume_directory=tmp_path,
        outreach_database_path=tmp_path / 'outreach.db', application_database_path=tmp_path / 'applications.db',
        notification_recipient_updater=None,
        company_config_path=Path(__file__).resolve().parents[1] / 'config/companies.yaml')
    return app, repo, ids


def test_detail_reuses_list_score_without_recomputing_and_invalidates_changed_job(workspace, monkeypatch):
    app, repo, ids = workspace
    matcher = Mock(wraps=web.match_resume_to_job)
    monkeypatch.setattr(web, 'match_resume_to_job', matcher)
    client = TestClient(app, base_url='http://127.0.0.1')
    page = client.get('/')
    assert page.status_code == 200 and matcher.call_count == 3
    detail = client.get(f'/jobs/{ids[0]}')
    assert matcher.call_count == 3
    expected = next(call.args for call in matcher.call_args_list if call.args[1].id == ids[0])
    result = matcher._mock_wraps(*expected)
    assert f'{result.score}/100' in detail.text
    assert result.eligibility_label in detail.text
    assert 'Not calculated' not in detail.text
    assert 'Refresh match' in detail.text
    job = repo.get_by_id(ids[0])
    job.requirements += ' Kubernetes required.'
    repo.save(job)
    assert client.get(f'/jobs/{ids[0]}').status_code == 200
    assert matcher.call_count == 4
    assert repo.get_by_id(ids[0]).status == 'new'


def test_changed_active_resume_invalidates_cached_score(workspace, tmp_path, monkeypatch):
    app, _, ids = workspace
    matcher = Mock(wraps=web.match_resume_to_job)
    monkeypatch.setattr(web, 'match_resume_to_job', matcher)
    client = TestClient(app, base_url='http://127.0.0.1')
    client.get(f'/jobs/{ids[0]}')
    client.get(f'/jobs/{ids[0]}')
    assert matcher.call_count == 1
    doc = Document()
    doc.add_paragraph('Demo Candidate\nSKILLS\nJava\nEDUCATION\nBSc Software Engineering in progress')
    content = BytesIO()
    doc.save(content)
    save_active_resume(tmp_path, 'updated-demo.docx', content.getvalue())
    assert client.get(f'/jobs/{ids[0]}').status_code == 200
    assert matcher.call_count == 2


def test_compact_cards_keep_every_blocker_and_escape_untrusted_explanations():
    macro = web.templates.env.get_template('_eligibility.html').module.eligibility
    reasons = ('Completed degree required. ' * 15, '<script>alert(1)</script> Two years experience required.')
    result = SimpleNamespace(eligibility_status='ineligible', eligibility_label='Ineligible',
                             eligibility_reasons=reasons)
    html = str(macro(result, True))
    assert reasons[0] in html
    assert 'Two years experience required.' in html
    assert '<script>' not in html and '&lt;script&gt;' in html


def test_no_active_resume_has_neutral_missing_information_state(workspace, tmp_path):
    app, _, ids = workspace
    (tmp_path / '.active_resume').unlink()
    # No matching file remains to trigger the existing resume fallback.
    for path in (tmp_path / 'uploaded').glob('*.docx'):
        path.unlink()
    client = TestClient(app, base_url='http://127.0.0.1')
    assert 'eligibility--unknown' in client.get('/').text
    assert 'Not calculated' in client.get(f'/jobs/{ids[0]}').text
