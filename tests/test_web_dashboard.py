from fastapi.testclient import TestClient
from datetime import datetime, timezone
from io import BytesIO
import re
from docx import Document

from src.models import JobPosting
from src.applications.service import discover_resume_pairs
from src.resume.resume_manager import save_active_resume
from src.storage.job_repository import JobRepository
from src.storage.postgres_job_repository import PostgresJobRepository
from src.web.app import _clean_captured_lines, _summarize_captured_lines, create_app


def test_postgres_dashboard_starts_without_schema_ownership(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("DATABASE_URL", "postgresql://dashboard@db.example/jobs")
    verified = []

    def forbidden_initialize(self):
        raise AssertionError("Dashboard must not run PostgreSQL migrations")

    monkeypatch.setattr(PostgresJobRepository, "initialize", forbidden_initialize)
    monkeypatch.setattr(PostgresJobRepository, "verify_schema", lambda self: verified.append(True))
    monkeypatch.setattr(PostgresJobRepository, "list_all", lambda self: [])
    client = TestClient(create_app(
        tmp_path / "unused.db",
        outreach_database_path=tmp_path / "outreach.db",
        application_database_path=tmp_path / "applications.db",
    ), base_url="http://127.0.0.1")
    assert verified == [True]
    assert client.get("/health").json()["storage"] == "postgres"


def _client_with_jobs(tmp_path):
    database = tmp_path / "jobs.db"
    repository = JobRepository(database)
    repository.initialize()
    job_id = repository.save(
        JobPosting(
            "Python Software Intern", "Example", "Remote, Israel", "https://example.com/jobs/1",
            description="Build FastAPI services.", requirements="Python and FastAPI",
        )
    )
    repository.save(
        JobPosting(
            "Unreviewed role", "Other", "Tel Aviv", "https://example.com/jobs/2",
            is_relevant=False,
        )
    )
    return TestClient(create_app(
        database, resume_directory=tmp_path,
        outreach_database_path=tmp_path / "outreach.db",
        application_database_path=tmp_path / "applications.db",
    ), base_url="http://127.0.0.1"), job_id


def test_outreach_requires_preview_and_approval_before_storing_draft(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    csrf = _csrf_token(client, job_id)
    added = client.post(f"/jobs/{job_id}/outreach/contacts", data={
        "csrf_token": csrf, "display_name": "Ari", "role_title": "Software Engineer",
        "company": "Example", "profile_url": "https://www.linkedin.com/in/ari?trk=test",
        "relationship": "hiring_team", "location": "Remote",
    })
    assert added.status_code == 200
    assert "Ari" in added.text
    assert "90/100" in added.text
    assert "Approved messages" not in added.text

    contact_id = re.search(r'name="contact_id" value="(\d+)"', added.text).group(1)
    preview = client.post(f"/jobs/{job_id}/outreach/preview", data={
        "csrf_token": _csrf_token(client, job_id), "contact_id": contact_id,
    })
    assert "Edit full message" in preview.text
    token = re.search(r'name="outreach_token" value="([^"]+)"', preview.text).group(1)
    approved = client.post(f"/jobs/{job_id}/outreach/approve", data={
        "csrf_token": _csrf_token(client, job_id), "outreach_token": token,
        "kind": "full_message", "body": "Edited and explicitly approved message.",
    })
    assert "Final approval" in approved.text
    confirmation = re.search(r'name="confirmation_token" value="([^"]+)"', approved.text).group(1)
    approved = client.post(f"/jobs/{job_id}/outreach/confirm", data={
        "csrf_token": _csrf_token(client, job_id), "confirmation_token": confirmation,
    })
    assert approved.status_code == 200
    assert "Approved messages" in approved.text
    assert "Edited and explicitly approved message." in approved.text
    assert "I sent this manually" in approved.text


def test_outreach_search_shortcuts_and_connection_note(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    page = client.get(f"/jobs/{job_id}")
    assert "Find Recruiters" in page.text
    assert "site%3Alinkedin.com" in page.text
    client.post(f"/jobs/{job_id}/outreach/contacts", data={
        "csrf_token": _csrf_token(client, job_id), "display_name": "Ari",
        "role_title": "Engineer", "company": "Example",
        "profile_url": "https://linkedin.com/in/ari", "relationship": "employee",
    })
    page = client.get(f"/jobs/{job_id}")
    contact_id = re.search(r'name="contact_id" value="(\d+)"', page.text).group(1)
    preview = client.post(f"/jobs/{job_id}/outreach/preview", data={
        "csrf_token": _csrf_token(client, job_id), "contact_id": contact_id,
        "kind": "connection_note",
    })
    assert 'name="kind" value="connection_note"' in preview.text
    body = re.search(r'<textarea name="body"[^>]*>(.*?)</textarea>', preview.text, re.S).group(1)
    assert len(body) <= 300


def test_outreach_tracks_manual_activity_without_sending(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    client.post(f"/jobs/{job_id}/outreach/contacts", data={
        "csrf_token": _csrf_token(client, job_id), "display_name": "Ari",
        "role_title": "Engineer", "company": "Example",
        "profile_url": "https://linkedin.com/in/ari", "relationship": "employee",
    })
    page = client.get(f"/jobs/{job_id}")
    contact_id = re.search(r'name="contact_id" value="(\d+)"', page.text).group(1)
    preview = client.post(f"/jobs/{job_id}/outreach/preview", data={
        "csrf_token": _csrf_token(client, job_id), "contact_id": contact_id,
        "kind": "full_message",
    })
    token = re.search(r'name="outreach_token" value="([^"]+)"', preview.text).group(1)
    approved = client.post(f"/jobs/{job_id}/outreach/approve", data={
        "csrf_token": _csrf_token(client, job_id), "outreach_token": token,
        "kind": "full_message", "body": "Approved body",
    })
    confirmation = re.search(r'name="confirmation_token" value="([^"]+)"', approved.text).group(1)
    approved = client.post(f"/jobs/{job_id}/outreach/confirm", data={
        "csrf_token": _csrf_token(client, job_id), "confirmation_token": confirmation,
    })
    draft_id = re.search(r'/outreach/drafts/(\d+)/activity', approved.text).group(1)
    sent = client.post(f"/jobs/{job_id}/outreach/drafts/{draft_id}/activity", data={
        "csrf_token": _csrf_token(client, job_id), "action": "sent",
        "follow_up_at": "2030-01-02",
    })
    assert "Sent manually" in sent.text
    assert 'value="2030-01-02"' in sent.text


def test_outreach_rejects_invalid_profile_url(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    response = client.post(f"/jobs/{job_id}/outreach/contacts", data={
        "csrf_token": _csrf_token(client, job_id), "display_name": "Ari",
        "role_title": "Engineer", "company": "Example", "profile_url": "javascript:bad",
    })
    assert "public HTTP(S) URL" in response.text
    assert "Tailored resume" in response.text
    assert "Set RESUME_DIRECTORY and create a tailored" not in response.text


def _csrf_token(client: TestClient, job_id: int) -> str:
    response = client.get(f"/jobs/{job_id}")
    match = re.search(r'name="csrf_token" value="([^"]+)"', response.text)
    assert match
    return match.group(1)


def _proposal_token(html: str) -> str:
    match = re.search(r'name="proposal_token" value="([^"]+)"', html)
    assert match
    return match.group(1)


def test_dashboard_filters_to_relevant_jobs_by_default(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)

    response = client.get("/")

    assert response.status_code == 200
    assert "Python Software Intern" in response.text
    assert "Unreviewed role" not in response.text
    assert "Opportunity desk" in response.text
    assert "AJob" in response.text
    assert "Job Agent" not in response.text
    assert "Relevant roles" in response.text
    assert "Needs review" in response.text
    assert 'aria-label="Filter jobs"' in response.text
    assert "Review opportunity" in response.text
    assert "company-generic.svg" in response.text
    assert "Update resume" in response.text
    assert 'enctype="multipart/form-data"' in response.text
    assert "Update email" in response.text
    assert 'action="/settings/notification-email"' in response.text


def test_dashboard_discloses_active_and_paused_scan_sources(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)

    response = client.get("/")

    assert "Scanning 48 companies" in response.text
    assert "Companies included in every scan" in response.text
    assert "Microsoft" in response.text
    assert "Applied Materials" in response.text
    assert "1 configured sources paused" in response.text
    assert "Palo Alto Networks / CyberArk (Idira)" in response.text
    assert "Check Point" in response.text
    assert "Wix" in response.text


def test_dashboard_reloads_scan_sources_from_current_config(tmp_path) -> None:
    config = tmp_path / "companies.yaml"
    config.write_text(
        "companies:\n  - name: First Source\n    careers_url: https://example.com/jobs\n",
        encoding="utf-8",
    )
    client = TestClient(create_app(
        tmp_path / "jobs.db",
        outreach_database_path=tmp_path / "outreach.db",
        application_database_path=tmp_path / "applications.db",
        company_config_path=config,
    ), base_url="http://127.0.0.1")
    assert "First Source" in client.get("/").text

    config.write_text(
        "companies:\n"
        "  - name: Second Source\n    careers_url: https://example.com/second\n"
        "  - name: Paused Source\n    careers_url: https://example.com/paused\n    enabled: false\n",
        encoding="utf-8",
    )
    response = client.get("/")

    assert "Scanning 1 companies" in response.text
    assert "Second Source" in response.text
    assert "Paused Source" in response.text
    assert "First Source" not in response.text


def test_dashboard_hides_jobs_ineligible_for_active_resume_by_default(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)
    repository = JobRepository(tmp_path / "jobs.db")
    repository.save(JobPosting(
        "Software Engineering Student", "Apple", "Jerusalem, Israel",
        "https://example.com/jobs/closed-major",
        description="Build software.",
        requirements=(
            "Minimum Qualifications: Pursuing a B.Sc./M.Sc. in Electrical Engineering, "
            "Computer Engineering."
        ),
    ))
    repository.save(JobPosting(
        "Software Engineering PhD Intern", "Google", "Tel Aviv, Israel",
        "https://example.com/jobs/phd-only",
        description="Build software.",
        requirements="Ph.D. in Computer Science required.",
    ))
    repository.save(JobPosting(
        "Junior Software Developer Intern", "Amazon", "Haifa, Israel",
        "https://example.com/jobs/poor-skill-fit",
        description="Build Ruby services.",
        requirements="Ruby required.",
    ))
    repository.save(JobPosting(
        "Python Developer Intern", "Microsoft", "Remote, Israel",
        "https://example.com/jobs/strong-skill-fit",
        description="Develop Python applications.",
        requirements="Python required.",
        discovered_at=datetime(2000, 1, 1, tzinfo=timezone.utc),
    ))
    resume = Document()
    resume.add_paragraph("Dana Developer")
    resume.add_paragraph("SKILLS\nPython")
    resume.add_paragraph(
        "EDUCATION\nB.Sc. Software Engineering, Example University\n2024-2027"
    )
    content = BytesIO()
    resume.save(content)
    save_active_resume(tmp_path, "resume.docx", content.getvalue())

    default_page = client.get("/")
    all_jobs_page = client.get("/", params={"relevance": "all"})

    assert "Python Software Intern" in default_page.text
    assert default_page.text.index("Python Developer Intern") < default_page.text.index(
        "Python Software Intern"
    )
    assert "Software Engineering Student" not in default_page.text
    assert "Software Engineering PhD Intern" not in default_page.text
    assert "Junior Software Developer Intern" not in default_page.text
    assert "Software Engineering Student" in all_jobs_page.text
    assert "Software Engineering PhD Intern" in all_jobs_page.text
    assert "Junior Software Developer Intern" in all_jobs_page.text
    assert "Not eligible:" in all_jobs_page.text


def test_student_eligibility_dashboard_possible_filter_and_status_preservation(tmp_path):
    client, _ = _client_with_jobs(tmp_path)
    repository = JobRepository(tmp_path / 'jobs.db')
    blocked_id = repository.save(JobPosting(
        'Graduate Bootcamp', 'Example', 'Israel', 'https://example.test/graduate',
        requirements='Qualifications\nB.Sc. in Software Engineering. Python.', status='applied'))
    repository.save(JobPosting('Junior Python Developer', 'Example', 'Israel', 'https://example.test/junior',
                               requirements='Python.'))
    document = Document()
    document.add_paragraph('Example Candidate\nThird-year Software Engineering student')
    document.add_paragraph('SKILLS\nPython')
    document.add_paragraph('EDUCATION\nB.Sc. Software Engineering (in progress)')
    content = BytesIO()
    document.save(content)
    save_active_resume(tmp_path, 'student.docx', content.getvalue())
    default = client.get('/')
    possible = client.get('/', params={'relevance': 'possible'})
    all_jobs = client.get('/', params={'relevance': 'all'})
    assert default.status_code == possible.status_code == all_jobs.status_code == 200
    assert 'Graduate Bootcamp' not in default.text
    assert default.text.index('Python Software Intern') < default.text.index('Junior Python Developer')
    assert 'Junior Python Developer' in possible.text
    assert 'Python Software Intern' not in possible.text
    assert 'Graduate Bootcamp' not in possible.text
    assert 'Potentially eligible' in possible.text
    assert 'Graduate Bootcamp' in all_jobs.text
    assert 'still studying' in all_jobs.text
    detail = client.post(f'/jobs/{blocked_id}/match', data={'csrf_token': _csrf_token(client, blocked_id)})
    assert 'Ineligible' in detail.text
    assert 'Technical fit:' in detail.text
    assert repository.get_by_id(blocked_id).status == 'applied'


def test_dashboard_reports_a_match_failure_for_only_the_affected_job(
    monkeypatch, tmp_path
) -> None:
    client, _ = _client_with_jobs(tmp_path)
    repository = JobRepository(tmp_path / "jobs.db")
    failed_job_id = repository.save(JobPosting(
        "Python Platform Intern", "Example", "Jerusalem, Israel",
        "https://example.com/jobs/match-error",
        description="Build Python services.",
        requirements="Python required.",
    ))
    resume = Document()
    resume.add_paragraph("SKILLS\nPython, FastAPI")
    resume.add_paragraph("EDUCATION\nB.Sc. Software Engineering\n2024-2027")
    content = BytesIO()
    resume.save(content)
    save_active_resume(tmp_path, "resume.docx", content.getvalue())

    import src.web.app as web_app

    original_match = web_app.match_resume_to_job

    def fail_one_posting(profile, job):
        if job.id == failed_job_id:
            raise ValueError("malformed posting")
        return original_match(profile, job)

    monkeypatch.setattr(web_app, "match_resume_to_job", fail_one_posting)
    default_page = client.get("/")
    all_jobs_page = client.get("/", params={"relevance": "all"})

    assert "Python Platform Intern" not in default_page.text
    assert "Match score unavailable for this job." in all_jobs_page.text


def test_dashboard_updates_cloud_notification_recipient_with_approval(tmp_path) -> None:
    updated: list[str] = []
    database = tmp_path / "jobs.db"
    JobRepository(database).initialize()
    client = TestClient(create_app(
        database,
        outreach_database_path=tmp_path / "outreach.db",
        application_database_path=tmp_path / "applications.db",
        notification_recipient_updater=updated.append,
    ), base_url="http://127.0.0.1")
    page = client.get("/")
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', page.text).group(1)

    response = client.post(
        "/settings/notification-email",
        data={"csrf_token": csrf, "notification_email": "new@example.com"},
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert updated == ["new@example.com"]
    assert "Daily email recipient updated" in response.text


def test_health_endpoint_identifies_current_ajob_server(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)

    response = client.get("/health")

    assert response.status_code == 200
    payload = response.json()
    assert payload["name"] == "AJob"
    assert payload["status"] == "ok"
    assert payload["storage"] == "sqlite"
    assert isinstance(payload["job_id"], int)
    assert datetime.fromisoformat(payload["started_at"]).tzinfo is not None
    assert client.get(f"/jobs/{payload['job_id']}").status_code == 200


def test_job_detail_uses_collapsible_content_sections(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)

    response = client.get(f"/jobs/{job_id}")

    assert response.status_code == 200
    assert response.text.count('class="content-section collapsible-section"') >= 5
    assert "<h2>About the role</h2>" in response.text
    assert "<h2>What they’re looking for</h2>" in response.text
    assert "<h2>Resume match</h2>" in response.text
    assert "<h2>Prepare application</h2>" in response.text
    assert "<h2>Contact people</h2>" in response.text
    assert "AJob" in response.text
    assert "company-generic.svg" in response.text


def test_job_detail_copy_removes_source_navigation_without_losing_job_facts() -> None:
    job = JobPosting(
        "Python Intern", "Example", "Tel Aviv", "https://example.com/job",
        description="Skip to main content\nApply now\nPython Intern\nBuild APIs.\nWork with mentors.",
    )

    assert _clean_captured_lines(job.description, job) == [
        "Build APIs.",
        "Work with mentors.",
    ]


def test_job_detail_copy_preserves_meaningful_leading_symbols_and_mixed_facts() -> None:
    job = JobPosting(
        "Systems Intern", "Example", "Tel Aviv", "https://example.com/job",
        description=(
            "Job ID: 12345\n"
            "Job ID: 12345 | Tel Aviv\n"
            "*nix experience\n"
            "-Werror knowledge\n"
            "- C++ experience"
        ),
    )

    assert _clean_captured_lines(job.description, job) == [
        "Job ID: 12345 | Tel Aviv",
        "*nix experience",
        "-Werror knowledge",
        "C++ experience",
    ]


def test_job_detail_summary_keeps_high_signal_sentences_and_limits_length() -> None:
    job = JobPosting(
        "Software Intern", "Example", "Tel Aviv", "https://example.com/job",
        description=(
            "Welcome to our growing organization. Build Python services with the team. "
            "Enjoy a collaborative office. Design reliable APIs. Develop internal tools. "
            "Work with product partners. Learn from mentors. Extra marketing copy."
        ),
    )

    result = _summarize_captured_lines(
        job.description, job,
        keywords=("build", "design", "develop", "team", "work with"),
        limit=4,
    )

    assert len(result) == 4
    assert "Build Python services with the team." in result
    assert "Design reliable APIs." in result
    assert "Welcome to our growing organization." not in result


def test_job_detail_summary_preserves_short_technical_requirements() -> None:
    job = JobPosting(
        "Software Intern", "Example", "Tel Aviv", "https://example.com/job",
        requirements="B.Sc.; C++; Python; Long promotional sentence without a requirement.",
    )

    result = _summarize_captured_lines(
        job.requirements, job,
        keywords=("b.sc", "c++", "python"),
        limit=3,
    )

    assert result == ["B.Sc.", "C++", "Python"]


def test_job_detail_exposes_full_source_when_one_line_contains_omitted_sentences(tmp_path) -> None:
    long_description = " ".join(f"Sentence {index} about the role." for index in range(1, 9))
    client, _ = _client_with_jobs(tmp_path)
    repository = JobRepository(tmp_path / "jobs.db")
    job_id = repository.save(
        JobPosting(
            "Software Intern", "Example", "Tel Aviv", "https://example.com/long-role",
            description=long_description,
        )
    )

    response = client.get(f"/jobs/{job_id}")

    assert response.status_code == 200
    assert "View full captured description" in response.text
    assert "Sentence 8 about the role." in response.text


def test_dashboard_uploads_resume_and_preserves_local_path_privacy(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)
    document_path = tmp_path / "upload-source.docx"
    document = Document()
    document.add_paragraph("Dana Developer\nSKILLS\nPython")
    document.save(document_path)
    page = client.get("/")
    token = re.search(r'name="csrf_token" value="([^"]+)"', page.text).group(1)

    response = client.post(
        "/resume",
        data={"csrf_token": token},
        files={
            "resume_file": (
                "Dana Resume.docx", document_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "Resume updated" in response.text
    assert "Dana Resume" in response.text
    assert "/100 match" in response.text
    assert "fit-suggestions" in response.text
    assert 'role="progressbar"' in response.text
    assert str(tmp_path) not in response.text
    assert (tmp_path / ".active_resume").is_file()


def test_resume_upload_rejects_invalid_csrf_without_writing(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)

    response = client.post(
        "/resume",
        data={"csrf_token": "wrong"},
        files={"resume_file": ("resume.pdf", b"%PDF-1.4", "application/pdf")},
    )

    assert response.status_code == 403
    assert not (tmp_path / ".active_resume").exists()


def test_resume_upload_rejects_request_without_content_length(tmp_path) -> None:
    client, _ = _client_with_jobs(tmp_path)
    request = client.build_request(
        "POST",
        "/resume",
        data={"csrf_token": "irrelevant"},
        files={"resume_file": ("resume.pdf", b"%PDF-1.4", "application/pdf")},
    )
    del request.headers["content-length"]

    response = client.send(request, follow_redirects=False)

    assert response.status_code == 303
    assert response.headers["location"] == "/?resume_error=invalid"
    assert not (tmp_path / ".active_resume").exists()


def test_dashboard_can_filter_and_update_a_job_status(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)

    filtered = client.get("/", params={"company": "Example", "query": "python"})
    updated = client.post(
        f"/jobs/{job_id}/status",
        data={"status": "applied", "csrf_token": _csrf_token(client, job_id)},
    )

    assert filtered.status_code == 200
    assert "Python Software Intern" in filtered.text
    assert updated.status_code == 200
    assert "Job status updated" in updated.text
    assert 'status--applied' in updated.text


def test_dashboard_hides_foreign_jobs_by_default_and_can_show_them(tmp_path) -> None:
    database = tmp_path / "jobs.db"
    repository = JobRepository(database)
    repository.initialize()
    repository.save(JobPosting(
        "Israel Software Intern", "Example", "Tel Aviv, ISR",
        "https://example.com/jobs/israel",
    ))
    repository.save(JobPosting(
        "Poland Software Intern", "Example", "Warsaw, Poland",
        "https://example.com/jobs/poland",
    ))
    client = TestClient(create_app(
        database, resume_directory=tmp_path,
        outreach_database_path=tmp_path / "outreach.db",
        application_database_path=tmp_path / "applications.db",
    ), base_url="http://127.0.0.1")

    default_page = client.get("/")
    all_locations = client.get("/", params={"location": "all"})

    assert "Israel Software Intern" in default_page.text
    assert "Poland Software Intern" not in default_page.text
    assert "Poland Software Intern" in all_locations.text


def test_job_detail_calculates_a_local_match_without_exposing_resume_text(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    resume = tmp_path / "resume.txt"
    resume.write_text(
        "Dana Developer\nSKILLS\nPython, FastAPI\nEDUCATION\nB.Sc. Computer Science",
        encoding="utf-8",
    )

    response = client.post(
        f"/jobs/{job_id}/match",
        data={"resume_path": str(resume), "csrf_token": _csrf_token(client, job_id)},
    )

    assert response.status_code == 200
    assert "/100 match" in response.text
    assert "Dana Developer" not in response.text
    assert "Set RESUME_DIRECTORY and create a tailored" not in response.text
    assert "Tailored resume" in response.text


def test_dashboard_previews_and_generates_an_explicitly_approved_tailored_resume(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    resume = tmp_path / "master_resume.txt"
    resume.write_text(
        "Dana Developer\nSKILLS\nPython, FastAPI\nEXPERIENCE\nBuilt a FastAPI service.\n"
        "EDUCATION\nB.Sc. Computer Science",
        encoding="utf-8",
    )
    preview = client.post(
        f"/jobs/{job_id}/match",
        data={"resume_path": str(resume), "csrf_token": _csrf_token(client, job_id)},
    )

    assert preview.status_code == 200
    assert "Tailor your resume" in preview.text
    assert "Current" in preview.text
    assert "After these edits" in preview.text
    assert "These edits improve presentation, not the score" in preview.text
    assert "Choose changes" in preview.text
    assert "Suggested" in preview.text
    assert "Why it’s supported" in preview.text

    generated = client.post(
        f"/jobs/{job_id}/tailor",
        data={
            "proposal_token": _proposal_token(preview.text),
            "approved_change_id": ["professional_format", "targeted_summary"],
            "csrf_token": _csrf_token(client, job_id),
        },
    )

    expected = tmp_path / "tailored" / "Example-Python-Software-Intern-tailored.docx"
    expected_pdf = tmp_path / "tailored" / "Example-Python-Software-Intern-tailored.pdf"
    assert generated.status_code == 200
    assert expected.is_file()
    assert expected_pdf.is_file()
    assert "master resume was not changed" in generated.text
    assert resume.read_text(encoding="utf-8").startswith("Dana Developer")


def test_dashboard_refuses_tailored_resume_without_explicit_approval(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    resume = tmp_path / "resume.txt"
    resume.write_text("Dana Developer\nSKILLS\nPython", encoding="utf-8")

    response = client.post(
        f"/jobs/{job_id}/tailor",
        data={"proposal_token": "missing", "csrf_token": _csrf_token(client, job_id)},
    )

    assert response.status_code == 200
    assert "preview expired" in response.text
    assert not (tmp_path / "tailored").exists()


def test_dashboard_rejects_approval_when_resume_changed_after_preview(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    resume = tmp_path / "resume.txt"
    resume.write_text("Dana Developer\nSKILLS\nPython", encoding="utf-8")
    preview = client.post(
        f"/jobs/{job_id}/match",
        data={"resume_path": str(resume), "csrf_token": _csrf_token(client, job_id)},
    )
    resume.write_text("Dana Developer\nSKILLS\nPython, FastAPI", encoding="utf-8")

    response = client.post(
        f"/jobs/{job_id}/tailor",
        data={
            "proposal_token": _proposal_token(preview.text),
            "approved_change_id": "professional_format",
            "csrf_token": _csrf_token(client, job_id),
        },
    )

    assert "changed after preview" in response.text
    assert not (tmp_path / "tailored").exists()


def _application_form(csrf: str, pdf, docx) -> dict[str, str]:
    return {
        "csrf_token": csrf,
        "full_name": "Dana Developer",
        "email": "dana@example.com",
        "phone": "+972 50-123-4567",
        "location": "Tel Aviv, Israel",
        "work_authorization": "authorized",
        "availability": "Immediately",
        "salary_expectation": "",
        "linkedin_url": "https://linkedin.com/in/dana",
        "github_url": "https://github.com/dana",
        "resume_pair": discover_resume_pairs(pdf.parent)[0].id,
        "cover_note": "I am interested in this role.",
    }


def _application_pair(tmp_path):
    from docx import Document
    from pypdf import PdfWriter

    tailored = tmp_path / "tailored"
    tailored.mkdir(exist_ok=True)
    pdf = tailored / "Example-Python-Software-Intern-tailored.pdf"
    docx = tailored / "Example-Python-Software-Intern-tailored.docx"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with pdf.open("wb") as output:
        writer.write(output)
    document = Document()
    document.add_paragraph("Resume")
    document.save(docx)
    return pdf, docx


def test_application_requires_exact_preview_then_records_manual_submission(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    pdf, docx = _application_pair(tmp_path)
    preview = client.post(
        f"/jobs/{job_id}/application/preview",
        data=_application_form(_csrf_token(client, job_id), pdf, docx),
    )

    assert preview.status_code == 200
    assert "Review and approve" in preview.text
    assert "does not submit an application" in preview.text
    token = re.search(r'name="application_token" value="([^"]+)"', preview.text).group(1)

    approved = client.post(
        f"/jobs/{job_id}/application/approve",
        data={"csrf_token": _csrf_token(client, job_id), "application_token": token},
    )
    assert approved.status_code == 200
    assert "Prepared packages" in approved.text
    assert "Open employer page manually" in approved.text
    package_id = re.search(r'/application/packages/(\d+)/submitted', approved.text).group(1)

    replay = client.post(
        f"/jobs/{job_id}/application/approve",
        data={"csrf_token": _csrf_token(client, job_id), "application_token": token},
    )
    assert "preview expired" in replay.text

    submitted = client.post(
        f"/jobs/{job_id}/application/packages/{package_id}/submitted",
        data={"csrf_token": _csrf_token(client, job_id), "confirm_submission": "yes"},
    )
    assert "Manual submission recorded" in submitted.text
    assert 'status--applied' in submitted.text


def test_application_form_lists_only_matching_tailored_resume_pairs(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    pdf, _ = _application_pair(tmp_path)
    unmatched = pdf.parent / "unmatched.pdf"
    unmatched.write_bytes(pdf.read_bytes())

    page = client.get(f"/jobs/{job_id}")

    assert f'>{pdf.stem}</option>' in page.text
    assert str(pdf.parent) not in page.text
    assert "unmatched</option>" not in page.text


def test_application_rejects_changed_resume_after_preview(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    pdf, docx = _application_pair(tmp_path)
    preview = client.post(
        f"/jobs/{job_id}/application/preview",
        data=_application_form(_csrf_token(client, job_id), pdf, docx),
    )
    token = re.search(r'name="application_token" value="([^"]+)"', preview.text).group(1)
    pdf.write_bytes(b"changed after review")

    response = client.post(
        f"/jobs/{job_id}/application/approve",
        data={"csrf_token": _csrf_token(client, job_id), "application_token": token},
    )

    assert "changed after preview" in response.text
    assert "Prepared packages" not in response.text


def test_application_rejects_changed_resume_when_recording_manual_submission(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    pdf, docx = _application_pair(tmp_path)
    preview = client.post(
        f"/jobs/{job_id}/application/preview",
        data=_application_form(_csrf_token(client, job_id), pdf, docx),
    )
    token = re.search(r'name="application_token" value="([^"]+)"', preview.text).group(1)
    approved = client.post(
        f"/jobs/{job_id}/application/approve",
        data={"csrf_token": _csrf_token(client, job_id), "application_token": token},
    )
    package_id = re.search(r'/application/packages/(\d+)/submitted', approved.text).group(1)
    pdf.write_bytes(b"changed after package approval")

    response = client.post(
        f"/jobs/{job_id}/application/packages/{package_id}/submitted",
        data={"csrf_token": _csrf_token(client, job_id), "confirm_submission": "yes"},
    )

    assert response.status_code == 200
    assert "changed after package approval" in response.text
    assert "Manual submission recorded" not in response.text
    assert JobRepository(tmp_path / "jobs.db").get_by_id(job_id).status == "new"


def test_application_rejects_unknown_or_malicious_resume_pair_id(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    outside_pdf = tmp_path / "resume.pdf"
    outside_docx = tmp_path / "resume.docx"
    outside_pdf.write_bytes(b"pdf")
    outside_docx.write_bytes(b"docx")

    response = client.post(
        f"/jobs/{job_id}/application/preview",
        data={**_application_form(_csrf_token(client, job_id), outside_pdf, outside_docx),
              "resume_pair": "../resume"},
    )

    assert "available tailored resume pair" in response.text
    assert "Review and approve" not in response.text


def test_application_rejects_pair_deleted_after_form_was_loaded(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    pdf, docx = _application_pair(tmp_path)
    form = _application_form(_csrf_token(client, job_id), pdf, docx)
    docx.unlink()

    response = client.post(f"/jobs/{job_id}/application/preview", data=form)

    assert "available tailored resume pair" in response.text
    assert "Review and approve" not in response.text


def test_application_preview_requires_csrf(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)
    pdf, docx = _application_pair(tmp_path)
    form = _application_form("", pdf, docx)

    response = client.post(f"/jobs/{job_id}/application/preview", data=form)

    assert response.status_code == 403


def test_dashboard_rejects_status_change_without_csrf_token(tmp_path) -> None:
    client, job_id = _client_with_jobs(tmp_path)

    response = client.post(f"/jobs/{job_id}/status", data={"status": "applied"})

    assert response.status_code == 403
import re
